from datetime import datetime
from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook

from app import create_app
from app.extensions import db
from app.models import Case, Customer, CustomerKind, Project, Task, User
from app.workflow import TaskPhase


def test_staff_monthly_statistics_scoped_to_assignee():
    app = create_app()
    suffix = uuid4().hex[:8]
    with app.app_context():
        staff_a = User(username=f"_staff_stat_a_{suffix}", role="staff")
        staff_a.set_password("secret")
        staff_b = User(username=f"_staff_stat_b_{suffix}", role="staff")
        staff_b.set_password("secret")
        customer = Customer(kind=CustomerKind.COMPANY, name=f"_staff_stat_customer_{suffix}")
        db.session.add_all([staff_a, staff_b, customer])
        db.session.flush()
        project = Project(customer_id=customer.id, name=f"_staff_stat_project_{suffix}")
        db.session.add(project)
        db.session.flush()

        case_a = Case(
            project_id=project.id,
            title=f"_staff_stat_mine_{suffix}",
            application_no=f"SSA{suffix}",
            case_type_code="invention_risk_precheck_mechanical",
            created_at=datetime(2088, 7, 10, 3, 0),
            actual_return_at=datetime(2088, 7, 16, 3, 0),
        )
        case_b = Case(
            project_id=project.id,
            title=f"_staff_stat_other_{suffix}",
            application_no=f"SSB{suffix}",
            case_type_code="utility_utility_model",
            created_at=datetime(2088, 7, 12, 3, 0),
            actual_return_at=datetime(2088, 7, 18, 3, 0),
        )
        db.session.add_all([case_a, case_b])
        db.session.flush()
        db.session.add_all(
            [
                Task(case_id=case_a.id, phase_status=TaskPhase.COMPLETED, assignee_id=staff_a.id),
                Task(case_id=case_b.id, phase_status=TaskPhase.COMPLETED, assignee_id=staff_b.id),
            ]
        )
        db.session.commit()
        staff_a_name = staff_a.username

    client = app.test_client()
    client.post("/auth/login", data={"username": staff_a_name, "password": "secret"}, follow_redirects=True)

    page = client.get("/staff/worklog?year=2088&month=7")
    assert page.status_code == 200
    assert "月度统计".encode() in page.data
    assert f"_staff_stat_mine_{suffix}".encode() not in page.data
    assert f"_staff_stat_other_{suffix}".encode() not in page.data
    assert b">1<" in page.data
    assert "发明".encode() in page.data
    assert "近 6 个月趋势".encode() in page.data
    assert page.data.count(b"qy-monthly-trend-point") == 6
    assert "2088年7月：1 件".encode() in page.data

    filtered = client.get(
        "/staff/case-detail?case_type_primary=invention&created_year=2088&created_month=7"
    )
    assert filtered.status_code == 200
    assert f"_staff_stat_mine_{suffix}".encode() in filtered.data
    assert f"_staff_stat_other_{suffix}".encode() not in filtered.data

    completed_page = client.get("/staff/worklog?year=2088&month=7&basis=completed")
    assert completed_page.status_code == 200
    assert "按实际返稿时间统计".encode() in completed_page.data

    exported = client.get("/staff/worklog/export?year=2088&month=7")
    assert exported.status_code == 200
    workbook = load_workbook(BytesIO(exported.data), read_only=True)
    rows = list(workbook.active.values)
    assert any(row[0] == "案件总数" and row[1] == 1 for row in rows)
    assert "案件明细" in workbook.sheetnames
    detail_rows = list(workbook["案件明细"].values)
    assert any(row[0] == f"_staff_stat_mine_{suffix}" and row[1] == f"SSA{suffix}" for row in detail_rows)
    assert all(row[0] != f"_staff_stat_other_{suffix}" for row in detail_rows[1:])
