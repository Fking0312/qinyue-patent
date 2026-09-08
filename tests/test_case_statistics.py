from datetime import datetime
from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook

from app import create_app
from app.extensions import db
from app.models import Case, Customer, CustomerKind, Project, User


def test_admin_case_statistics_month_type_filter_and_export():
    app = create_app()
    suffix = uuid4().hex[:8]
    with app.app_context():
        admin = User(username=f"_stats_admin_{suffix}", role="admin")
        admin.set_password("secret")
        staff = User(username=f"_stats_staff_{suffix}", role="staff")
        staff.set_password("secret")
        customer = Customer(kind=CustomerKind.COMPANY, name=f"_stats_customer_{suffix}")
        db.session.add_all([admin, staff, customer])
        db.session.flush()
        project = Project(customer_id=customer.id, name=f"_stats_project_{suffix}")
        db.session.add(project)
        db.session.flush()
        cases = [
            Case(
                project_id=project.id,
                title=f"_stats_invention_{suffix}",
                application_no=f"STI{suffix}",
                case_type_code="invention_risk_precheck_mechanical",
                created_at=datetime(2088, 7, 10, 3, 0),
                actual_return_at=datetime(2088, 7, 16, 3, 0),
            ),
            Case(
                project_id=project.id,
                title=f"_stats_utility_{suffix}",
                application_no=f"STU{suffix}",
                case_type_code="utility_utility_model",
                created_at=datetime(2088, 7, 20, 3, 0),
                actual_return_at=datetime(2088, 6, 28, 3, 0),
            ),
            Case(
                project_id=project.id,
                title=f"_stats_previous_{suffix}",
                application_no=f"STP{suffix}",
                case_type_code="trademark_trademark",
                created_at=datetime(2088, 6, 20, 3, 0),
            ),
        ]
        db.session.add_all(cases)
        db.session.commit()
        admin_name = admin.username
        staff_name = staff.username

    client = app.test_client()
    client.post("/auth/login", data={"username": admin_name, "password": "secret"})

    page = client.get("/admin/flow-map?year=2088&month=7")
    assert page.status_code == 200
    assert "案件统计".encode() in page.data
    assert "发明".encode() in page.data
    assert "实用新型".encode() in page.data
    assert b">2<" in page.data
    assert b"created_year=2088" in page.data

    filtered = client.get(
        "/admin/cases?case_type_primary=invention&created_year=2088&created_month=7"
    )
    assert filtered.status_code == 200
    assert f"_stats_invention_{suffix}".encode() in filtered.data
    assert f"_stats_utility_{suffix}".encode() not in filtered.data
    assert f"_stats_previous_{suffix}".encode() not in filtered.data

    exported = client.get("/admin/flow-map/export?year=2088&month=7")
    assert exported.status_code == 200
    assert exported.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    workbook = load_workbook(BytesIO(exported.data), read_only=True)
    rows = list(workbook.active.values)
    assert ("案件总数", 2, None, None, None) in rows
    assert any(row[0] == "发明" and row[1] == 1 for row in rows)

    completed_page = client.get("/admin/flow-map?year=2088&month=7&basis=completed")
    assert completed_page.status_code == 200
    assert "按实际返稿时间统计".encode() in completed_page.data
    assert b"actual_return_year=2088" in completed_page.data

    completed_filtered = client.get(
        "/admin/cases?case_type_primary=invention"
        "&actual_return_year=2088&actual_return_month=7"
    )
    assert completed_filtered.status_code == 200
    assert f"_stats_invention_{suffix}".encode() in completed_filtered.data
    assert f"_stats_utility_{suffix}".encode() not in completed_filtered.data

    completed_export = client.get(
        "/admin/flow-map/export?year=2088&month=7&basis=completed"
    )
    assert completed_export.status_code == 200
    completed_workbook = load_workbook(BytesIO(completed_export.data), read_only=True)
    completed_rows = list(completed_workbook.active.values)
    assert any(row[0] == "统计口径" and row[1] == "实际返稿时间" for row in completed_rows)
    assert any(row[0] == "案件总数" and row[1] == 1 for row in completed_rows)

    client.get("/auth/logout")
    client.post("/auth/login", data={"username": staff_name, "password": "secret"})
    forbidden = client.get("/admin/flow-map")
    assert forbidden.status_code == 403
